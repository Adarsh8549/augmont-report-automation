"""
parse_docx.py — Extract report sections from Augmont daily .docx file.
Outputs a JSON object to stdout.

Usage:
    python scripts/parse_docx.py reports/report.docx
"""

import sys
import json
import re
from pathlib import Path

try:
    from docx import Document
except ImportError:
    print("Installing python-docx...", file=sys.stderr)
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx", "-q"])
    from docx import Document


def extract_date_from_filename(path: str) -> str:
    name = Path(path).stem
    match = re.search(r'(\d{1,2})\s*(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\s*(\d{4})', name, re.IGNORECASE)
    if match:
        day, month, year = match.groups()
        return f"{month.capitalize()} {day}, {year}"
    return ""


def strip_prefix(text: str, *prefixes) -> str:
    for prefix in prefixes:
        if text.lower().startswith(prefix.lower()):
            return text[len(prefix):].strip(" –-:")
    return text


def parse_support_resistance(table) -> dict:
    result = {}
    full_text = "\n".join(cell.text for row in table.rows for cell in row.cells)

    def find_value(pattern):
        m = re.search(pattern, full_text, re.IGNORECASE)
        return m.group(1).strip() if m else ""

    result["gold_intl_support"]      = find_value(r"International Gold Support Level\s*[:\-]\s*(\$[\d,./]+)")
    result["gold_intl_resistance"]   = find_value(r"International Gold Resistance Level\s*[:\-]\s*(\$[\d,./]+)")
    result["gold_dom_support"]       = find_value(r"Domestic Gold Support Level\s*[:\-]\s*(Rs\s*[\d,./]+\s*/?\s*\d*\s*gm?)")
    result["gold_dom_resistance"]    = find_value(r"Domestic Gold Resistance Level\s*[:\-]\s*(Rs\s*[\d,./]+\s*/?\s*\d*\s*gm?)")
    result["silver_intl_support"]    = find_value(r"International Silver Support Level\s*[:\-]\s*(\$[\d,./]+)")
    result["silver_intl_resistance"] = find_value(r"International Silver Resistance Level\s*[:\-]\s*(\$[\d,./]+)")
    result["silver_dom_support"]     = find_value(r"Domestic Silver Support Level\s*[:\-]\s*(Rs\s*[\d,./]+\s*/?\s*kg?)")
    result["silver_dom_resistance"]  = find_value(r"Domestic Silver Resistance Level\s*[:\-]\s*(Rs\s*[\d,./]+\s*/?\s*kg?)")

    # Pre-formatted strings ready for Canva replace_text
    result["canva_gold_intl_support"]      = f"Gold Support Level: {result['gold_intl_support']}"
    result["canva_gold_intl_resistance"]   = f"Gold Resistance Level: {result['gold_intl_resistance']}"
    result["canva_gold_dom_support"]       = f"Gold Support Level: {result['gold_dom_support']}"
    result["canva_gold_dom_resistance"]    = f"Gold Resistance Level: {result['gold_dom_resistance']}"
    result["canva_silver_intl_support"]    = f"Silver Support Level: {result['silver_intl_support']}"
    result["canva_silver_intl_resistance"] = f"Silver Resistance Level: {result['silver_intl_resistance']}"
    result["canva_silver_dom_support"]     = f"Silver Support Level: {result['silver_dom_support']}"
    result["canva_silver_dom_resistance"]  = f"Silver Resistance Level: {result['silver_dom_resistance']}"
    return result


def parse(docx_path: str) -> dict:
    doc = Document(docx_path)
    data = {
        "headline": "",
        "date": extract_date_from_filename(docx_path),
        "price_movement": "",
        "geopolitical": "",
        "macro": "",
        "gold_technicals": "",
        "silver_technicals": "",
    }

    paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    skip = {"fundamental news and triggers", "technical triggers", "support and resistance"}

    for text in paras:
        tl = text.lower()
        if tl in skip or tl.startswith("disclaimer"):
            continue
        if tl.startswith("price movement"):
            data["price_movement"] = strip_prefix(text, "Price Movement")
        elif tl.startswith("geopolitical"):
            data["geopolitical"] = strip_prefix(text, "Geopolitical Developments", "Geopolitical")
        elif tl.startswith("macro"):
            data["macro"] = strip_prefix(text, "Macro-economic Signals", "Macro-economic", "Macro")
        elif tl.startswith("gold") and not data["gold_technicals"] and "support" not in tl:
            data["gold_technicals"] = text
        elif tl.startswith("silver") and not data["silver_technicals"] and "support" not in tl:
            data["silver_technicals"] = text
        elif not data["headline"]:
            data["headline"] = text

    if doc.tables:
        data.update(parse_support_resistance(doc.tables[0]))

    first_sentence = data["price_movement"].split(".")[0].strip()
    data["excerpt"] = first_sentence + "." if first_sentence else ""
    return data


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python parse_docx.py <path_to_docx>", file=sys.stderr)
        sys.exit(1)
    path = sys.argv[1]
    if not Path(path).exists():
        print(f"Error: file not found: {path}", file=sys.stderr)
        sys.exit(1)
    print(json.dumps(parse(path), indent=2, ensure_ascii=False))
